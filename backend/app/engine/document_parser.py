import io
import re
from typing import List, Dict, Any
import docx

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False


class DocumentQuestionParser:
    """
    Parses questions and extracts raw contextual text from uploaded documents (.pdf, .docx, .txt, .csv).
    Automatically extracts question text, detects marks, and isolates units.
    """

    @classmethod
    def extract_raw_text(cls, file_bytes: bytes, filename: str) -> str:
        """
        Extracts clean plain text content from PDF, DOCX, or TXT file bytes.
        """
        lower_name = filename.lower()
        if lower_name.endswith('.pdf'):
            return cls.extract_text_from_pdf(file_bytes)
        elif lower_name.endswith('.docx'):
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        else:
            try:
                return file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                return file_bytes.decode('latin-1', errors='ignore')

    @classmethod
    def extract_text_from_pdf(cls, file_bytes: bytes) -> str:
        if not PYPDF_AVAILABLE:
            raise RuntimeError("PDF parser engine (pypdf) is not installed.")
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(text.strip())
        return "\n\n".join(pages_text)

    @classmethod
    def parse_txt(cls, content: str) -> List[Dict[str, Any]]:
        lines = content.splitlines()
        questions = []
        current_unit = 1

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                continue

            # Detect Unit Header (e.g., "UNIT 1", "UNIT-2", "Unit III")
            unit_match = re.search(r'\bunit\s*[-:]?\s*([1-5]|I{1,3}|IV|V)\b', trimmed, re.IGNORECASE)
            if unit_match:
                roman_map = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}
                val = unit_match.group(1).upper()
                current_unit = roman_map.get(val, int(val) if val.isdigit() else current_unit)
                continue

            # Pipe-delimited row: Unit | Marks | Question | Difficulty
            if '|' in trimmed:
                parts = [p.strip() for p in trimmed.split('|')]
                if len(parts) >= 3:
                    questions.append({
                        "unit_number": int(parts[0]) if parts[0].isdigit() else current_unit,
                        "marks": int(parts[1]) if parts[1].isdigit() else 2,
                        "question_text": parts[2],
                        "difficulty": parts[3].upper() if len(parts) > 3 and parts[3].upper() in ['EASY', 'MEDIUM', 'HARD'] else 'MEDIUM'
                    })
                    continue

            # Numbered question pattern (e.g. "1. Define ... [2M]", "Q12: Explain ... (10 Marks)")
            cleaned = re.sub(r'^(?:Q\d+|\d+)[\.\:\)]\s*', '', trimmed)
            
            # Detect marks in brackets like [2M] or (10 Marks)
            marks_match = re.search(r'[\(\[]\s*(\d+)\s*(?:m|marks|mark)?\s*[\)\]]', cleaned, re.IGNORECASE)
            marks = int(marks_match.group(1)) if marks_match else (2 if len(cleaned.split()) < 15 else 10)
            cleaned = re.sub(r'[\(\[]\s*\d+\s*(?:m|marks|mark)?\s*[\)\]]', '', cleaned).strip()

            if len(cleaned) >= 8:
                questions.append({
                    "unit_number": current_unit,
                    "marks": marks,
                    "question_text": cleaned,
                    "difficulty": "HARD" if marks >= 10 else "MEDIUM"
                })

        return questions

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        raw_text = cls.extract_raw_text(file_bytes, "doc.docx")
        return cls.parse_txt(raw_text)

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        raw_text = cls.extract_text_from_pdf(file_bytes)
        return cls.parse_txt(raw_text)

    @classmethod
    def parse_any(cls, file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
        lower = filename.lower()
        if lower.endswith('.pdf'):
            return cls.parse_pdf(file_bytes)
        elif lower.endswith('.docx'):
            return cls.parse_docx(file_bytes)
        else:
            try:
                txt = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                txt = file_bytes.decode('latin-1', errors='ignore')
            return cls.parse_txt(txt)
